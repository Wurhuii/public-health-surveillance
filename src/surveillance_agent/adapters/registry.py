from __future__ import annotations

from typing import List, Optional, Type

from .base import Adapter, read_rows
from .care_home import CareHomeAdapter, HospitalAdapter, SchoolAdapter

ADAPTERS: List[Type[Adapter]] = [
    HospitalAdapter,
    CareHomeAdapter,
    SchoolAdapter,
]


def detect_adapter(header: List[str]) -> Optional[Type[Adapter]]:
    for cls in ADAPTERS:
        if cls.recognizes(header):
            return cls
    return None


def detect_adapter_from_file(path: str) -> Optional[Type[Adapter]]:
    if not path.lower().endswith((".xlsx", ".xlsm")):
        import csv
        try:
            with open(path, "r", encoding="utf-8-sig", errors="replace", newline="") as fh:
                reader = csv.DictReader(fh)
                header = reader.fieldnames or []
        except Exception:
            return None
    else:
        try:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            ws = wb.active
            header = [str(c) for c in next(ws.iter_rows(values_only=True), [])]
        except Exception:
            return None
    return detect_adapter([h or "" for h in header])
