from __future__ import annotations

import csv
import os
from typing import Any, Dict, Iterator, List, Optional

from ..domain import SurveillanceEvent
from ..utils import hash_id, normalize_text


def read_rows(path: str) -> Iterator[Dict[str, str]]:
    if not os.path.isfile(path):
        return
    if path.lower().endswith((".xlsx", ".xlsm")):
        yield from _read_xlsx(path)
    else:
        with open(path, "r", encoding="utf-8-sig", errors="replace", newline="") as fh:
            reader = csv.DictReader(fh)
            for i, row in enumerate(reader):
                d = {k: normalize_text(v) for k, v in row.items()}
                d["_index"] = str(i)
                yield d


def _read_xlsx(path: str) -> Iterator[Dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = [normalize_text(c) for c in next(rows, [])]
    for i, values in enumerate(rows):
        d = {header[j]: normalize_text(values[j]) for j in range(len(header))}
        d["_index"] = str(i)
        yield d


class Adapter:
    source: str = ""

    @classmethod
    def recognizes(cls, header: List[str]) -> bool:
        raise NotImplementedError

    def parse(self, row: Dict[str, str], salt: str) -> Optional[SurveillanceEvent]:
        raise NotImplementedError

    def quality_issues(self, row: Dict[str, str]) -> List[str]:
        return []
